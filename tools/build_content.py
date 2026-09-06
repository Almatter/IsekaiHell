from docx import Document
from docx.oxml.ns import qn
from pathlib import Path
from PIL import Image
import io, json, re, shutil, sys

ROOT=Path(__file__).resolve().parents[1]
RACES_DOC=Path(sys.argv[1]) if len(sys.argv)>1 else Path('/mnt/data/Races.docx')
LORE_DOC=Path(sys.argv[2]) if len(sys.argv)>2 else Path('/mnt/data/World Lore.docx')

TREE_NAMES=['Prime','Beast','Fae','Monster','Construct','Hybrid']
KNOWN_FIELDS={'Lifespan','Average Height','Average Weight','Average Physique','Body Tint, Colouring and Marking','Historical Figures','Interspecies Relations and Assumptions','Pre-req:'}
COUNTRIES=[
 ('ryke','Ryke',['Ryke','Isekai Hell - Ryke']),
 ('east','Otenzian Empire (East)',['Otenzian Empire (East)','Isekai Hell - Otenzian Empire (East)']),
 ('republic','The Republic',['The Republic','Isekai Hell - The Republic']),
 ('fae-see','See of Chearon',['See of Chearon','Isekai Hell - See of Fae']),
 ('west','Ororoot Empire (West)',['Ororoot Empire (West)','Isekai Hell - West Empire']),
 ('lake','Continental Lake',['Continental Lake','Isekai Hell - Continental Lake']),
 ('duchy','The Grand Duchy',['The Grand Duchy','Isekai Hell - The Grand Duchy']),
 ('kingdom','Kingdom of Rotia',['Kingdom of Rotia','Isekai Hell - Kingdom']),
 ('sky','Sky Nation',['Sky Nation','Isekai Hell - Sky Nation']),
 ('widersia','Widersia',['Widersia','Isekai Hell - Widersia']),
]
THEMES={
 'ryke':{'accent':'#d7b45a','accent2':'#b12d39','glow':'#72c1b6','deep':'#18090c','motif':'trade'},
 'east':{'accent':'#e0a34c','accent2':'#a72f22','glow':'#f0c36e','deep':'#1d0e08','motif':'imperial'},
 'republic':{'accent':'#8fc06e','accent2':'#3f7447','glow':'#c4d89f','deep':'#09140c','motif':'wild'},
 'fae-see':{'accent':'#75d7b7','accent2':'#6c4cab','glow':'#a8f0dd','deep':'#071411','motif':'fae'},
 'west':{'accent':'#d4b45d','accent2':'#8b5b25','glow':'#91c18a','deep':'#171008','motif':'desert'},
 'lake':{'accent':'#76cde8','accent2':'#2c6e9c','glow':'#b6ecff','deep':'#06121a','motif':'water'},
 'duchy':{'accent':'#b59af1','accent2':'#664fb2','glow':'#d6c8ff','deep':'#100a1b','motif':'arcane'},
 'kingdom':{'accent':'#e2c46c','accent2':'#395e9e','glow':'#dce9ff','deep':'#080e18','motif':'crown'},
 'sky':{'accent':'#d8ecff','accent2':'#6ca4cf','glow':'#ffffff','deep':'#09131a','motif':'sky'},
 'widersia':{'accent':'#c99355','accent2':'#616c74','glow':'#efd0a4','deep':'#12100e','motif':'gear'},
}

def clean(s):
    return re.sub(r'[ \t\r\f\v]+',' ',(s or '').replace('\u200b','').replace('\u200e','').replace('\ufeff','')).strip()

def norm(s):
    return re.sub(r'[^a-z0-9]+','',clean(s).lower())

def slug(s):
    v=re.sub(r'[^a-z0-9]+','-',clean(s).lower()).strip('-')
    return v or 'entry'

def para_images(p):
    out=[]
    for run in p.runs:
        for blip in run._r.xpath('.//a:blip'):
            rid=blip.get(qn('r:embed'))
            if rid and rid in p.part.related_parts:
                part=p.part.related_parts[rid]
                out.append((str(part.partname).lstrip('/'),part.blob))
    return out

def save_webp(blob,out_path,max_w=720,max_h=900,quality=82):
    try:
        im=Image.open(io.BytesIO(blob))
        try: im.seek(0)
        except Exception: pass
        if im.mode not in ('RGB','RGBA'):
            im=im.convert('RGBA' if 'transparency' in im.info else 'RGB')
        if im.mode=='RGBA':
            bg=Image.new('RGBA',im.size,(0,0,0,0)); bg.alpha_composite(im); im=bg
        im.thumbnail((max_w,max_h),Image.Resampling.LANCZOS)
        out_path.parent.mkdir(parents=True,exist_ok=True)
        im.save(out_path,'WEBP',quality=quality,method=6)
        return True
    except Exception:
        return False

def split_lines(text):
    lines=[]
    for raw in (text or '').split('\n'):
        x=clean(raw)
        if x: lines.append(x)
    return lines

def race_extract():
    doc=Document(RACES_DOC)
    paras=doc.paragraphs
    actual=[]
    current_tree=''
    current_stage='Mundane'
    for i,p in enumerate(paras):
        t=clean(p.text); st=p.style.name
        if st=='Title' and t in TREE_NAMES:
            current_tree=t; current_stage='Mundane'
        elif st=='Heading 1' and t:
            m=re.search(r'\b(Mundane|Intermediate|Ascended)\b',t,re.I)
            if m: current_stage=m.group(1).title()
        elif st=='Heading 2' and t:
            m=re.match(r'^(Mundane|Intermediate|Ascended)\s+',t,re.I)
            if m:
                current_stage=m.group(1).title(); continue
            if t in ['Variant Path','Native Path'] or (current_tree=='Hybrid' and t in ['Prime','Beast','Fae','Monster','Construct']):
                continue
            # look ahead, ignoring blank styled paras often used only to hold images
            h3=[]; j=i+1
            while j<len(paras):
                q=paras[j]; qt=clean(q.text); qst=q.style.name
                if qt and qst in ['Heading 2','Heading 1','Title']:
                    break
                if qt and qst=='Heading 3': h3.append(qt)
                j+=1
            if any(x in KNOWN_FIELDS for x in h3):
                actual.append((i,t,current_tree,current_stage,j))
    out=[]
    assets=ROOT/'assets'/'races'
    if assets.exists(): shutil.rmtree(assets)
    assets.mkdir(parents=True,exist_ok=True)
    for idx,(start,name,tree,stage,end) in enumerate(actual):
        desc=[]; fields={}; current_field=None; images=[]
        for p in paras[start+1:end]:
            t=clean(p.text); st=p.style.name
            for _,blob in para_images(p):
                images.append(blob)
            if t and st=='Heading 3':
                current_field=t; fields.setdefault(current_field,[]); continue
            if not t: continue
            vals=split_lines(p.text)
            if current_field: fields.setdefault(current_field,[]).extend(vals)
            else: desc.extend(vals)
        description=' '.join(desc)
        description=re.sub(r'^Description:\s*','',description,flags=re.I)
        pre=fields.get('Pre-req:',[])
        # pull inline heading-ish labels from pre-req into convenient metadata
        evolves=''; total=''
        for line in pre:
            if re.match(r'^Evolves into:',line,re.I): evolves=line.split(':',1)[1].strip()
            if re.match(r'^Total Cost:',line,re.I): total=line.split(':',1)[1].strip()
        img_paths=[]
        for k,blob in enumerate(images[:2],1):
            rel=f"assets/races/{slug(tree)}-{slug(stage)}-{slug(name)}-{k}.webp"
            if save_webp(blob,ROOT/rel,max_w=520,max_h=720,quality=80): img_paths.append(rel)
        aliases={name}
        for line in pre:
            if re.match(r'^Titles?:',line,re.I):
                for token in re.split(r'[,/]',line.split(':',1)[1]):
                    title=token.strip().strip('[]')
                    if title and title not in TREE_NAMES: aliases.add(title)
        if name=='Elf (Lune)': aliases.add('Moon Elf')
        entry={
            'id':f"{slug(tree)}-{slug(stage)}-{slug(name)}",
            'name':name,'tree':tree,'stage':stage,'description':description,
            'lifespan':' '.join(fields.get('Lifespan',[])),
            'height':' '.join(fields.get('Average Height',[])),
            'weight':' '.join(fields.get('Average Weight',[])),
            'physique':' '.join(fields.get('Average Physique',[])),
            'coloring':' '.join(fields.get('Body Tint, Colouring and Marking',[])),
            'historicalFigures':fields.get('Historical Figures',[]),
            'relations':fields.get('Interspecies Relations and Assumptions',[]),
            'requirements':pre,'evolvesInto':evolves,'totalCost':total,'images':img_paths,
            'aliases':sorted(aliases)
        }
        out.append(entry)
    js='window.IH_RACES = '+json.dumps(out,ensure_ascii=False,separators=(',',':'))+';\n'
    (ROOT/'data'/'races.js').write_text(js,encoding='utf-8')
    return out

def lore_blocks(paras,start,end,asset_prefix,max_images=7):
    blocks=[]; image_no=0
    for p in paras[start:end]:
        t=clean(p.text); st=p.style.name
        imgs=para_images(p)
        for _,blob in imgs:
            if image_no>=max_images: break
            image_no+=1
            rel=f"assets/lore/{asset_prefix}-{image_no:02d}.webp"
            if save_webp(blob,ROOT/rel,max_w=1100,max_h=850,quality=80):
                blocks.append({'type':'image','src':rel,'alt':f'{asset_prefix} lore image {image_no}'})
        if not t: continue
        if st=='Title': typ='h1'
        elif st=='Heading 1': typ='h2'
        elif st=='Heading 2': typ='h3'
        elif st=='Heading 3': typ='h4'
        elif st=='Subtitle': typ='subtitle'
        else: typ='p'
        blocks.append({'type':typ,'text':'\n'.join(split_lines(p.text))})
    return blocks

def lore_extract():
    doc=Document(LORE_DOC); paras=doc.paragraphs
    assets=ROOT/'assets'/'lore'
    if assets.exists(): shutil.rmtree(assets)
    assets.mkdir(parents=True,exist_ok=True)
    title_positions=[]
    for i,p in enumerate(paras):
        t=clean(p.text)
        if p.style.name=='Title' and t: title_positions.append((i,t))
    # main country starts by exact first-title names
    starts={}
    for cid,display,aliases in COUNTRIES:
        for i,t in title_positions:
            if norm(t)==norm(aliases[0]): starts[cid]=i; break
    world_start=next(i for i,t in title_positions if norm(t)==norm('World'))
    first_country=min(starts.values())
    intro=lore_blocks(paras,world_start+1,first_country,'world',max_images=6)
    countries=[]
    ordered=sorted([(starts[cid],cid,display,aliases) for cid,display,aliases in COUNTRIES])
    for pos,(start,cid,display,aliases) in enumerate(ordered):
        end=ordered[pos+1][0] if pos+1<len(ordered) else next((i for i,t in title_positions if i>start and norm(t)==norm('RP List')),len(paras))
        blocks=lore_blocks(paras,start,end,cid,max_images=7)
        # summary = first meaningful paragraph after titles/headings
        summary=''
        for b in blocks:
            if b['type']=='p' and len(b.get('text',''))>90:
                summary=b['text'][:420].strip(); break
        countries.append({'id':cid,'name':display,'aliases':aliases,'summary':summary,'theme':THEMES[cid],'blocks':blocks})
    payload={'intro':intro,'countries':countries,'source':'World Lore.docx'}
    (ROOT/'data'/'lore.js').write_text('window.IH_LORE = '+json.dumps(payload,ensure_ascii=False,separators=(',',':'))+';\n',encoding='utf-8')
    return payload

if __name__=='__main__':
    races=race_extract(); lore=lore_extract()
    print(f'Wrote {len(races)} race templates and {len(lore["countries"])} country lore sections.')
